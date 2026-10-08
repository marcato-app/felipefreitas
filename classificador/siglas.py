"""Dicionário de siglas de laboratório (3 letras) para o fim do descritivo de Farmácia: "NIMESULIDA 100MG 12CPRS (GER)".

Fontes (sem duplicar): marcas da DIMA com o código no fim ("NIMESULIDA GER", "DORFLEX (OPE)"), códigos da base de marcas de
Farmácia do cliente (Hoja, farmaCodLab) e a base de VITAMINA E MINERAL (marca "SUNDOWN NES", descrição "... (NES)").
- codigos: sigla -> fabricante (o que mais aparece com ela, >= 50% e 2+ vezes)
- fab: fabricante -> sigla padrão (a que tem as letras do nome: HYPERA -> HYP; senão a mais usada); fabricante sem
  nenhuma sigla nas bases ganha uma gerada (3 primeiras letras do nome, sem repetir), marcada para revisão
Saída: consts.json -> farmaSiglas e dados/farma_siglas_laboratorios.xlsx (para conferir).
Uso: python3 siglas.py && python3 build.py   (depois de dima.py e farma_marcas.py)
"""
import csv, gzip, json, re, unicodedata
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill

here = Path(__file__).parent
N = lambda s: ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().upper()).split())
SEM = {'OUTRO FABRICANTE', 'SIN PROVEEDOR ASOCIADO', 'SIN PROVEEDOR', 'POR DEFECTO', ''}
LIXO = set('LABORATORIO LABORATORIOS LAB LABS FARMACEUTICA FARMACEUTICOS FARMACEUTICO FARMA PHARMA PHARMACEUTICA IND '
           'INDUSTRIA COM COMERCIO SA S A LTDA ME EIRELI DO DA DE E BRASIL BR CIA GRUPO'.split())
NAO_SIGLA = set('MARCA OUTRA OUTRO COM SEM CPR COMP CAP CAPS CPS DRG GTS SOL SUS XPE AMP INJ CRE POM GEL ADT INF PED UND '
                'MCG MEQ REV OPC GOT KIT MAX PRO NEW ONE DAY MIX FIT TOP KID VIT MAG ZERO SUN MET DUO FOR XR RET ORO ZIN CAL FER OMG'.split())  # palavra de produto (MET = metformina)
# palavras que não distinguem um fabricante de outro (razão social): somem na chave do nome único
FAB_LIXO = LIXO | set('LABORATORIOS FARMACETICA FARMACEUTICAS FARMACEUTICOS FARMACEUTICO IMPORTACAO EXPORTACAO IMP EXP MEDICAMENTOS DISTRIBUIDORA'.split())
chave_fab = lambda f: ' '.join(w for w in N(f).split() if w not in FAB_LIXO)
ok_sigla = lambda s: bool(re.fullmatch(r'(?=.*[A-Z])[A-Z0-9]{3}', s)) and s not in NAO_SIGLA  # 1FA (1FARMA) também vale


def main():
    c = json.loads((here / 'consts.json').read_text(encoding='utf-8'))
    cod = defaultdict(Counter)  # sigla -> fabricante -> ocorrências
    def junta(marca_ou_desc, fab):
        if N(fab) in SEM: return
        m = re.search(r'\(([A-Z0-9]{3})\)\s*$', str(marca_ou_desc).upper()) or re.match(r'^.*\S\s+([A-Z0-9]{3})$', N(marca_ou_desc))
        if m and ok_sigla(m.group(1)): cod[m.group(1)][fab.strip()] += 1
    d = json.loads(gzip.decompress((here / 'dima.json.gz').read_bytes()))
    r = d['rows']
    for k in range(0, len(r), 3): junta(d['marcas'][r[k]], d['fabs'][r[k + 2]])
    for s, f in (c.get('farmaCodLab') or {}).items():
        if ok_sigla(s): cod[s][f] += 3  # base de marcas de Farmácia do cliente
    for p in sorted((here / 'dados' / 'vitaminas').glob('*.tsv')):
        for b in csv.DictReader(open(p, encoding='utf-8'), delimiter='\t'):
            junta(b['Marca'], b['Fabricante']); junta(b['Descripcion'], b['Fabricante'])
    codigos = {}
    for s, fs in cod.items():
        f, v = fs.most_common(1)[0]
        if v / sum(fs.values()) >= 0.5 and v >= 2: codigos[s] = f
    # fabricante -> sigla padrão
    por_fab = defaultdict(Counter)
    for s, f in codigos.items(): por_fab[f][s] += cod[s][f]
    fab, rel = {}, []
    iniciais = lambda f: ''.join(w for w in N(f).split() if w not in LIXO)
    for f, ss in por_fab.items():
        ini = iniciais(f)
        pref = [s for s in ss if ini.startswith(s)]
        fab[N(f)] = pref[0] if pref else ss.most_common(1)[0][0]
    # código que a base Hoja usa para o laboratório vale como sigla padrão dele, mesmo sendo palavra comum (FER = FERRING)
    for sg, f in (c.get('farmaCodLab') or {}).items():
        if re.fullmatch(r'(?=.*[A-Z])[A-Z0-9]{3}', sg) and N(f) not in fab: fab[N(f)] = sg
    # fabricantes das bases sem nenhuma sigla: gera (3 primeiras letras do nome, sem colidir com sigla de outro)
    # só fabricantes de Farmácia: EST MER 6 com código ATC ou de categoria da cesta Farmácia, base Hoja e base de vitaminas
    farma = {N(x['nome']) for x in c['libSeed']['categorias'] + c['cfg']['categorias'] if x.get('cesta') == 'FARMACIA'} | {'VITAMINA E MINERAL'}
    seg_farma = lambda sg: bool(re.match(r'^[A-Z]\d\d', sg)) or N(sg) in farma
    todos = Counter(d['fabs'][r[k + 2]] for k in range(0, len(r), 3) if seg_farma(d['segs'][r[k + 1]]) and N(d['fabs'][r[k + 2]]) not in SEM)
    for p in sorted((here / 'dados' / 'vitaminas').glob('*.tsv')):
        for b in csv.DictReader(open(p, encoding='utf-8'), delimiter='\t'):
            if N(b['Fabricante']) not in SEM: todos[b['Fabricante'].strip()] += 1
    for m in c.get('farmaMarcas', []):
        if m[2] and N(m[2]) not in SEM: todos[m[2]] += 1
    usadas = set(codigos)
    geradas = {}
    for f in sorted(todos):
        if N(f) in fab: continue
        ini = iniciais(f) or N(f).replace(' ', '')
        cands = [ini[:3]] + [ini[0] + x + y for x in ini[1:] for y in ini[2:]] if len(ini) >= 3 else []
        s = next((x for x in cands if ok_sigla(x) and x not in usadas), None)
        if s: fab[N(f)] = s; usadas.add(s); geradas[N(f)] = s
    # sigla de cada marca como a base de marcas de Farmácia do cliente (Hoja) escreve: NATZ RDF, PICOPREP FER, VIVACITA PRO
    # (vale antes da sigla padrão do fabricante; inclui siglas que são palavra comum, só para aquela marca)
    marcas = {}
    try:
        from farma_marcas import ler
        for pz in sorted((here / 'dados' / 'farma_marcas').glob('*')):
            if pz.suffix.lower() not in ('.tsv', '.csv', '.txt'): continue
            for f, m, _e7, v in ler(pz):
                x = re.match(r'^(.*\S)\s+\(?([A-Z0-9]{3})\)?$', N(m) if '(' not in str(m) else str(m).upper().strip())
                if x and re.fullmatch(r'(?=.*[A-Z])[A-Z0-9]{3}', x.group(2)): marcas.setdefault(N(x.group(1)), Counter())[x.group(2)] += (v or 1)
    except Exception as e:
        print('Hoja:', e)
    # nome genérico com vários laboratórios (VITAMINA D3 BIO / UNI): sem sigla fixa, vale a do fabricante do produto
    marcas = {k: cn.most_common(1)[0][0] for k, cn in marcas.items() if len(cn) == 1}
    # nome único de cada fabricante: grafias do mesmo fabricante (SANOFI FARMACEUTICA / SANOFI FARMACEUTICA LTDA,
    # GROSS / LABORATORIO GROSS) viram UM nome que já existe. Prioridade da grafia: 1º DIMA (mais confiável),
    # 2º base Hoja (tem erros de digitação), 3º base VITAMINA E MINERAL; dentro da mesma base, a mais usada
    fonte = [Counter(), Counter(), Counter()]
    for k in range(0, len(r), 3):
        f = d['fabs'][r[k + 2]]
        if seg_farma(d['segs'][r[k + 1]]) and N(f) not in SEM: fonte[0][f.strip()] += 1
    for m in c.get('farmaMarcas', []):
        if m[2] and N(m[2]) not in SEM: fonte[1][m[2].strip()] += 1
    for p in sorted((here / 'dados' / 'vitaminas').glob('*.tsv')):
        for b in csv.DictReader(open(p, encoding='utf-8'), delimiter='\t'):
            if N(b['Fabricante']) not in SEM: fonte[2][b['Fabricante'].strip()] += 1
    canon, nivel, pk = {}, {}, Counter()
    for i, fc in enumerate(fonte):
        for f, n in sorted(fc.items(), key=lambda x: (-x[1], len(x[0]))):
            k = chave_fab(f); pk[k] += n
            if k and k not in canon: canon[k] = f; nivel[k] = i
    # erro de digitação da Hoja/vitaminas (PRATI DONADUZI): 1 letra de diferença de um nome da DIMA -> o nome da DIMA
    def lev1(x, y):
        if abs(len(x) - len(y)) != 1: return False  # só letra/espaço a mais ou a menos (VILLAGE x SILLAGE é outro)
        i = 0
        while i < min(len(x), len(y)) and x[i] == y[i]: i += 1
        return x[i:] == y[i + 1:] or x[i + 1:] == y[i:]
    dk = [k for k in canon if nivel[k] == 0]
    typo = []
    for k in [k for k in canon if nivel[k] > 0 and len(k) >= 6]:
        cand = [x for x in dk if len(x) >= 6 and not x.isdigit() and lev1(k, x)]
        if len(cand) == 1: typo.append((canon[k], canon[cand[0]])); canon[k] = canon[cand[0]]
    print('erros de digitação corrigidos pela DIMA:', typo[:15], len(typo))
    # nome curto que é começo de nome mais completo: vai para o completo quando ele domina (UNIAO -> UNIAO QUIMICA,
    # FORHEALTH -> FORHEALTH NUTRICIONAL); VITA, APIS (vários nomes começam assim, nenhum domina) não juntam
    juntou = []
    for k in list(canon):
        if ' ' in k or len(k) < 4 or N(canon[k]) != k: continue  # GLOBAL MEDICAMENTOS, SOUL BRASIL: a palavra que sumiu pode distinguir
        longos = sorted((x for x in canon if x.startswith(k + ' ')), key=lambda x: (nivel[x], -pk[x]))
        if not longos: continue
        top = max(longos, key=lambda x: pk[x]); tot = sum(pk[x] for x in longos)
        if pk[top] >= 3 * pk[k] and pk[top] >= 0.8 * tot and nivel[top] <= nivel[k]:
            juntou.append((canon[k], canon[top])); canon[k] = canon[top]
    print('juntados pelo começo do nome:', juntou)
    # razão social da CMED que não está nas bases de Farmácia: o fabricante do Hoja completo (todas as cestas) com o
    # mesmo nome vale antes de criar nome novo (VASCONCELOS FARMACEUTICA E -> VASCONCELOS)
    hf = here / 'dados' / 'hoja_fabricantes_todos.txt'
    cmed_hoja = []
    if hf.exists():
        hk = defaultdict(set)
        for f in open(hf, encoding='utf-8'):
            f = f.strip(); k = chave_fab(f)
            if k: hk[k].add(f)
        labs = set()
        for p in sorted((here / 'dados' / 'cmed').glob('*.xlsx')):
            for row in openpyxl.load_workbook(p, read_only=True).active.iter_rows(values_only=True):
                if len(row) > 5 and row[2] and str(row[5] or '').strip()[:1].isdigit(): labs.add(str(row[2]).strip())
        tira = lambda f: re.sub(r'\b(LTDA|S A|SA|EIRELI|ME|EPP|INDUSTRIA|COMERCIO|IND|COM|E|DE|DO|DA|DOS|DAS)\b', ' ', N(f))
        GEN = set(('PRODUTOS PRODUTO QUIMICA QUIMICAS QUIMICOS QUIMICO FARMACO HOSPITALARES HOSPITALAR MEDICOS MEDICO MEDICAS GERAIS '
                   'IMPORTADORA EXPORTADORA DISTRIBUICAO FORNECIMENTO COMERCIAL NACIONAL BRASILEIRA PESQUISA CIENTIFICA SERVICOS '
                   'EQUIPAMENTOS MATERIAIS COSMETICOS ALIMENTOS NATURAIS FUNDACAO INSTITUTO EMPRESA CENTRO ESTADO').split())
        for k in [k for k in hk if len(k) < 4 or all(x in GEN for x in k.split())]: del hk[k]  # PRODUTOS DA DA, DO GERAIS
        for lab in sorted(labs):
            if re.search(r'\b(ESTADO|GOVERNADOR|UNIVERSIDADE|COMANDO)\b', N(lab)): continue  # órgão público: nome de lugar não é fabricante do Hoja
            w = chave_fab(tira(lab)).split()
            while len(w) > 1 and w[0] in GEN: w.pop(0)  # PRODUTOS ROCHE QUIMICOS -> ROCHE
            if not w or w[0] in GEN: continue
            if any(' '.join(w[:n]) in canon and (n == len(w) or n >= 2 or len(' '.join(w[:n])) >= 5) for n in range(len(w), 0, -1)): continue
            for n in range(len(w), 0, -1):
                k = ' '.join(w[:n])
                if not (n == len(w) or n >= 2 or len(k) >= 5): continue
                if len(hk.get(k, ())) == 1:
                    canon[' '.join(w)] = next(iter(hk[k])); cmed_hoja.append((lab, canon[' '.join(w)])); break
    print('CMED -> fabricante do Hoja completo:', len(cmed_hoja), cmed_hoja[:20])
    c['farmaFabs'] = canon
    c['farmaSiglas'] = {'codigos': codigos, 'fab': fab, 'marcas': marcas}
    (here / 'consts.json').write_text(json.dumps(c, ensure_ascii=False), encoding='utf-8')
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Siglas'
    ws.append(['SIGLA', 'FABRICANTE', 'OCORRENCIAS NAS BASES', 'ORIGEM', 'SIGLA PADRAO DO FABRICANTE?'])
    for x in ws[1]: x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F3864')
    for s in sorted(codigos): ws.append([s, codigos[s], sum(cod[s].values()), 'bases (DIMA/Hoja/vitaminas)', 'SIM' if fab.get(N(codigos[s])) == s else ''])
    for f, s in sorted(geradas.items(), key=lambda x: x[1]): ws.append([s, f, 0, 'GERADA (conferir)', 'SIM'])
    for col, w in zip('ABCDE', [8, 45, 12, 28, 14]): ws.column_dimensions[col].width = w
    ws.auto_filter.ref = ws.dimensions; ws.freeze_panes = 'A2'
    wb.save(here / 'dados' / 'farma_siglas_laboratorios.xlsx')
    print(f"{len(canon)} fabricantes (nome único; {len(set().union(*fonte)) - len(set(canon.values()))} grafias juntadas)")
    print(f"{len(marcas)} marcas com sigla na Hoja; {len(codigos)} siglas das bases ({len(por_fab)} fabricantes), {len(geradas)} geradas; "
          f"ex.: EMS PHARMA={fab.get('EMS PHARMA')} HYPERA={fab.get('HYPERA PHARMA')} GER={codigos.get('GER')} -> dados/farma_siglas_laboratorios.xlsx")


if __name__ == '__main__':
    main()
