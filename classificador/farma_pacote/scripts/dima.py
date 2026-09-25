"""Converte dados/DIMA_Peso_Fixo.xlsx na base de marcas embutida no app (dima.json.gz -> classificador.html).

Abas:
- "Mapeio marca-categoria-fabricante": EST_MER_6_DESCRIPCION, MARCA, FABRICANTE  -> marca x categoria x fabricante
- "Abreviaturas": MARCA, Abr_1..Abr_n  -> outras formas da marca aparecer na descrição
- "Sheet3": categoria, OUTRA MARCA, OUTRO FABRICANTE -> padrão de cada categoria quando não há marca
  (é o que o app já faz; guardado só como referência, não trava marca)

O app procura a marca primeiro nesta base; só vai para o dicionário enviado se não achar nada aqui.
Uso: python3 dima.py && python3 build.py
"""
import gzip, json
from pathlib import Path
import openpyxl

here = Path(__file__).parent
XLSX = here / 'dados' / 'DIMA_Peso_Fixo.xlsx'


def main():
    ws = openpyxl.load_workbook(XLSX, read_only=True).worksheets
    s = lambda v: str(v).strip() if v is not None else ''
    segs, fabs, marcas = {}, {}, {}
    idx = lambda d, k: d.setdefault(k, len(d))
    rows = []
    for r in ws[0].iter_rows(min_row=2, values_only=True):
        seg, mar, fab = s(r[0]), s(r[1]), s(r[2])
        if not mar: continue
        rows += [idx(marcas, mar), idx(segs, seg), idx(fabs, fab or 'OUTRO FABRICANTE')]
    # Farmácia: a DIMA grava a marca com o código do laboratório ("DORFLEX (OPE)", "NEOSALDINA HYP"), mas a descrição
    # traz só "DORFLEX". Nas categorias ATC, o nome sem o código também vira marca, com as mesmas linhas.
    import re
    atc = re.compile(r'^[A-Z]\d\d[A-Z]?\d?\s')
    inv_m = {v: k for k, v in marcas.items()}; inv_s = {v: k for k, v in segs.items()}
    # só marca comercial: nome com princípio ativo (genérico: "BICARBONATO DE SODIO XXX", "OMEPRAZOL TEU") fica de fora
    import unicodedata
    Nn = lambda t: ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().upper()).split())
    pa = sorted({Nn(a) for r in json.loads((here / 'consts.json').read_text(encoding='utf-8')).get('farmaPA', [])
                 for p in r['partes'] for a in p.split('|') if Nn(a)}, key=len, reverse=True)
    rx_pa = re.compile(r'(?<![A-Z0-9])(?:' + '|'.join(map(re.escape, pa)) + r')(?![A-Z0-9])') if pa else None
    extra = []
    for k in range(0, len(rows), 3):
        mar, seg = inv_m[rows[k]], inv_s[rows[k + 1]]
        if not atc.match(seg + ' '): continue
        m = re.match(r'^(.*?)\s*\(([A-Z0-9]{2,4})\)$', mar) or re.match(r'^(.*\S)\s+([A-Z0-9]{3})$', mar)
        if not m: continue
        base = m.group(1).strip()
        if len(base) >= 3 and not base.isdigit() and not (rx_pa and rx_pa.search(Nn(base))): extra += [idx(marcas, base), rows[k + 1], rows[k + 2]]
    rows += extra
    abr = []
    for r in ws[1].iter_rows(min_row=2, values_only=True):
        mar = s(r[0])
        if not mar: continue
        a = [s(x) for x in r[1:] if s(x)]
        abr.append([idx(marcas, mar)] + a)  # marca sem abreviação também entra (lista de marcas válidas)
    sem = sorted({s(r[0]) for r in ws[2].iter_rows(values_only=True) if s(r[0]) and s(r[1]).upper() == 'OUTRA MARCA'})
    inv = lambda d: [k for k, _ in sorted(d.items(), key=lambda kv: kv[1])]
    out = {'segs': inv(segs), 'fabs': inv(fabs), 'marcas': inv(marcas), 'rows': rows, 'abr': abr, 'semMarca': sem}
    raw = json.dumps(out, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    (here / 'dima.json.gz').write_bytes(gzip.compress(raw, 9))
    print(f"{len(marcas)} marcas, {len(fabs)} fabricantes, {len(segs)} categorias, {len(rows)//3} linhas, "
          f"{sum(len(a)-1 for a in abr)} abreviações, {len(sem)} categorias sem marca; "
          f"{len(raw)/1e6:.1f} MB -> {(here / 'dima.json.gz').stat().st_size/1e6:.1f} MB gzip")


if __name__ == '__main__':
    main()
