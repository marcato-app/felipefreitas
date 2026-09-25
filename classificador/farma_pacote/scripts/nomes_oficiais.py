"""Nome oficial (EST MER 7 / EST MER 6 com o código) das categorias cujo nome interno perdeu o código.

Uma limpeza antiga tirou o código ATC do começo do nome de ~380 categorias de Farmácia
("A02B1 ANTAGONISTAS H2" virou "ANTAGONISTAS H2"). O nome interno fica como está (regras, ligações e o que está
salvo no navegador dependem dele); o nome oficial vai em consts.json -> nomesOficiais e é o que sai na planilha e
o que é reconhecido na coluna CATEGORIA do arquivo. Fontes: dicionário de princípios ativos (Est Mer 6/7) e DIMA.
Nomes sem fonte ficam como estão até chegar a árvore oficial (pôr em dados/arvore_est_mer.xlsx: colunas
EST MER 7 / EST MER 6 com código e rodar de novo).

Uso: python3 nomes_oficiais.py && python3 build.py
"""
import gzip, json, re, unicodedata
from pathlib import Path
import openpyxl

here = Path(__file__).parent
N = lambda s: ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper()).split())
COD = re.compile(r'^[A-Z]\d{2}[A-Z]?\d?\s+')
tira = lambda s: COD.sub('', s.strip())


# decididos com o cliente quando a fonte tem mais de um código para o mesmo nome
FIXOS = {'CORTICOIDES': 'S01B CORTICOIDES'}  # R03D ou S01B: cliente escolheu o oftalmológico


def conserta(nome):
    """mojibake (ANÃ\\x81LOGOS -> ANALOGOS) e hífen que sobrou no lugar do código"""
    try: t = nome.encode('latin-1').decode('utf-8')
    except (UnicodeEncodeError, UnicodeDecodeError): t = nome
    t = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode()
    return re.sub(r'^[-\s]+', '', t).strip()


def main():
    c = json.loads((here / 'consts.json').read_text(encoding='utf-8'))
    cats = c['libSeed']['categorias'] + c['cfg']['categorias']
    fonte = {}
    def junta(v):
        v = str(v).strip()
        if v and N(v) != 'POR DEFECTO' and COD.match(v): fonte.setdefault(N(tira(v)), set()).add(v)
    for r in openpyxl.load_workbook(here / 'dados' / 'dicionario_nomenclaturas_farmaceuticas.xlsx', read_only=True)['Dicionário'].iter_rows(min_row=2, values_only=True):
        for v in (r[4], r[5]):
            if v: junta(v)
    for s in json.loads(gzip.decompress((here / 'dima.json.gz').read_bytes()))['segs']: junta(s)
    arv = here / 'dados' / 'arvore_est_mer.xlsx'
    if arv.exists():
        for ws in openpyxl.load_workbook(arv, read_only=True).worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v: junta(v)
    out, amb = {}, []
    for x in cats:
        k = N(conserta(x['nome'])); f = {v for v in fonte.get(k, set()) if ' - ' not in v} or fonte.get(k, set())
        cod = lambda v: v.split(' ')[0]
        f = {v for v in f if not any(o != v and cod(o).startswith(cod(v)) for o in f)}  # código mais específico
        if len(f) == 1: out[x['nome']] = next(iter(f))
        elif len(f) > 1: amb.append((x['nome'], sorted(f)))
        elif conserta(x['nome']) != x['nome']: out[x['nome']] = conserta(x['nome'])
    out.update({k: v for k, v in FIXOS.items() if any(x['nome'] == k for x in cats)})
    amb = [a for a in amb if a[0] not in FIXOS]
    c['nomesOficiais'] = out
    (here / 'consts.json').write_text(json.dumps(c, ensure_ascii=False), encoding='utf-8')
    far = [x['nome'] for x in cats if x.get('cesta') == 'FARMACIA']
    print(f"{len(out)} nomes oficiais ({sum(1 for n in far if n in out)} de {len(far)} de Farmácia); ambíguos: {amb}")


if __name__ == '__main__':
    main()
