"""Lista de preços da CMED/Anvisa (PMC - xls) -> cmed.json.gz, embutido no app pelo build.py (/*CMED*/null).

Por código de barras (EAN 1/2/3): PRODUTO, APRESENTAÇÃO (quantidade oficial), LABORATÓRIO, SUBSTÂNCIA,
CLASSE TERAPÊUTICA (EphMRA: "L1H - INIBIDORES DE PROTEINA QUINASE") e TIPO (Genérico, Similar, Novo...).
A planilha tem ~40 linhas de aviso antes do cabeçalho; o script acha a linha com APRESENTAÇÃO e EAN.
Usa o arquivo mais novo de dados/cmed/. Uso: python3 cmed.py [--baixar] [--historico] && python3 build.py
  --baixar: baixa do site da Anvisa a lista PMC mais nova e os dados abertos de medicamentos registrados.
  --historico: baixa as edições anteriores (página "anos anteriores") para dados/cmed/historico/ (resumo compacto).
  Códigos de barras que só existem em edições antigas (produto que saiu de linha) entram com a data da última edição;
  a edição mais nova sempre vence.

Medicamentos registrados (dados abertos, DADOS_ABERTOS_MEDICAMENTOS.csv): nome do produto -> princípio ativo e
categoria regulatória (Genérico, Similar, Novo...). Entra só o que não está na lista de preços; a categoria vem do
princípio ativo (dicionário do cliente no app, ou a classe que a substância tem na CMED).
"""
import csv, gzip, io, json, re, sys, unicodedata, urllib.request
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl

here = Path(__file__).parent
N = lambda s: ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper()).split())


PAGINA = 'https://www.gov.br/anvisa/pt-br/assuntos/medicamentos/cmed/precos'
PAGINA_DCB = 'https://www.gov.br/anvisa/pt-br/assuntos/farmacopeia/dcb'
DADOS = 'https://dados.anvisa.gov.br/dados/DADOS_ABERTOS_MEDICAMENTOS.csv'


def baixa(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=120) as r: return r.read()


def baixar():
    d = here / 'dados' / 'cmed'; d.mkdir(parents=True, exist_ok=True)
    html = baixa(PAGINA).decode('utf-8', 'ignore')
    m = re.search(r'href="([^"]*xls_conformidade_site_(\d{8})[^"]*)"', html)  # "PMC - xls"
    if m:
        alvo = d / f'cmed_pmc_{m.group(2)}.xlsx'
        if not alvo.exists():
            url = m.group(1) if m.group(1).startswith('http') else 'https://www.gov.br' + m.group(1)
            alvo.write_bytes(baixa(url)); print('baixado', alvo.name)
        else: print('lista PMC já é a mais nova:', alvo.name)
    (d / 'DADOS_ABERTOS_MEDICAMENTOS.csv').write_bytes(baixa(DADOS)); print('baixado DADOS_ABERTOS_MEDICAMENTOS.csv')
    # DCB: lista consolidada mais nova ("3-2024-lista-consolidada-dcb-ago.xlsx")
    h = baixa(PAGINA_DCB).decode('utf-8', 'ignore').replace('\\u002F', '/')
    ls = set(re.findall(r'(https://www\.gov\.br/anvisa/pt-br/assuntos/farmacopeia/dcb/(\d)-(\d{4})-lista-consolidada-dcb[\w-]*\.xlsx)', h))
    if ls:
        url, n, ano = max(ls, key=lambda x: (x[2], x[1]))
        alvo = d / f'dcb_{ano}_{n}.xlsx'
        if not alvo.exists():
            for velho in d.glob('dcb_*.xlsx'): velho.unlink()
            alvo.write_bytes(baixa(url + '/@@download/file')); print('baixado', alvo.name)


def registrados(cprod, sub_cls):
    """nome do produto -> [princípio ativo, classe CMED pela substância, categoria regulatória] (só o que não está na CMED)"""
    f = here / 'dados' / 'cmed' / 'DADOS_ABERTOS_MEDICAMENTOS.csv'
    if not f.exists(): return []
    rows = csv.DictReader(io.StringIO(f.read_bytes().decode('latin-1')), delimiter=';')
    por = defaultdict(Counter)
    for r in rows:
        nome = N(r.get('NOME_PRODUTO', ''))
        if not nome or nome in cprod or len(nome) < 4 or 'DINAMIZADO' in N(r.get('CATEGORIA_REGULATORIA', '')): continue
        pa = ';'.join(N(x) for x in re.split(r'[;,]', r.get('PRINCIPIO_ATIVO', '')) if N(x))
        if not pa: continue
        por[nome][(pa, N(r.get('CATEGORIA_REGULATORIA', '')))] += 2 if r.get('SITUACAO_REGISTRO') == 'Ativo' else 1
    out = []
    for nome, c in por.items():
        (pa, tipo), _ = c.most_common(1)[0]
        subs = sorted(x.strip() for x in pa.split(';') if x.strip())
        cls = sub_cls.get(';'.join(subs)) or (sub_cls.get(subs[0]) if len(subs) == 1 else '') or ''
        out.append([nome, ' + '.join(subs), cls, tipo])
    return out


def ler_planilha(path):
    """linhas da lista PMC (xls antigo ou xlsx): [eans], produto, apresentação, laboratório, classe, substância, tipo"""
    if str(path).lower().endswith('.xls'):
        import xlrd
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        it = (sh.row_values(i) for i in range(sh.nrows))
    else:
        it = openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True)
    H = None
    for r in it:
        if r and any('APRESENTA' in str(v or '').upper() for v in r) and any('EAN' in str(v or '').upper() for v in r):
            H = [N(v or '').replace(' ', '') for v in r]; break
    if not H: return []
    col = lambda *ts: next((i for t in ts for i, x in enumerate(H) if x.startswith(t)), -1)
    iS, iL, iA, iC, iT = col('SUBSTANCIA', 'PRINCIPIOATIVO'), col('LABORATORIO'), col('APRESENTACAO'), col('CLASSETERAPEUTICA'), col('TIPODEPRODUTO')
    iP = H.index('PRODUTO') if 'PRODUTO' in H else col('PRODUTO')
    iE = [i for i, x in enumerate(H) if x.startswith('EAN')]
    txt = lambda r, i: str(r[i] if i >= 0 and i < len(r) and r[i] is not None else '').strip()
    out = []
    for r in it:
        if not r or len(r) <= iP or not r[iP]: continue
        eans = []
        for i in iE:
            v = r[i] if i < len(r) else ''
            if isinstance(v, float): v = '%.0f' % v
            e = re.sub(r'\D', '', str(v or '')).lstrip('0')
            if len(e) >= 8: eans.append(e)
        if eans: out.append([eans, txt(r, iP), txt(r, iA), txt(r, iL), txt(r, iC), txt(r, iS), txt(r, iT)])
    return out


def historico():
    """baixa as edições anteriores da PMC (página "anos anteriores") e guarda um resumo compacto de cada uma"""
    d = here / 'dados' / 'cmed' / 'historico'; d.mkdir(parents=True, exist_ok=True)
    html = baixa(PAGINA + '/anos-anteriores').decode('utf-8', 'ignore')
    links = sorted(set(re.findall(r'href="([^"]*xls_conformidade_site_(\d{8})[^"]*)"', html)), key=lambda x: x[1])
    for url, data_ in links:
        alvo = d / f'{data_}.json.gz'
        if alvo.exists(): continue
        url = url if url.startswith('http') else 'https://www.gov.br' + url
        ext = '.xlsx' if '.xlsx' in url else '.xls'
        tmp = d / f'_tmp{ext}'
        try:
            tmp.write_bytes(baixa(url)); linhas = ler_planilha(tmp)
            alvo.write_bytes(gzip.compress(json.dumps(linhas, ensure_ascii=False, separators=(',', ':')).encode('utf-8'), 9))
            print(data_, len(linhas), 'linhas', flush=True)
        except Exception as e: print(data_, 'ERRO', e, flush=True)
        finally:
            if tmp.exists(): tmp.unlink()


def dcb(sub_cls):
    """DCB (Denominação Comum Brasileira): sais ("CLORIDRATO DE X" -> X) e substância -> classe CMED"""
    arqs = sorted((here / 'dados' / 'cmed').glob('dcb_*.xlsx'))
    if not arqs: return [], []
    nomes = set()
    for r in openpyxl.load_workbook(arqs[-1], read_only=True).worksheets[0].iter_rows(values_only=True):
        if len(r) > 3 and r[1] and str(r[3] or '').strip() in ('INF', 'IFA', 'BIO', 'PM'): nomes.add(N(r[1]))
    sais, base_de = Counter(), {}
    for n in nomes:
        m = re.match(r'^(\w+) DE (.+)$', n)
        if m and m.group(2) in nomes: sais[m.group(1)] += 1; base_de[n] = m.group(2)
    # classe por substância base (produto de uma substância só na CMED, com ou sem o sal)
    cls = defaultdict(Counter)
    for k, c in sub_cls.items():
        if ';' in k or not k: continue
        b = base_de.get(k, k)
        if b in nomes: cls[b][c] += 1
    subs = [[b, c.most_common(1)[0][0]] for b, c in sorted(cls.items()) if len(b) >= 5]
    return subs, sorted(w for w, n in sais.items() if n >= 2)


def main():
    if '--baixar' in sys.argv: baixar()
    if '--historico' in sys.argv: historico()
    arqs = sorted((here / 'dados' / 'cmed').glob('cmed_pmc_*.xls*'))
    if not arqs: print('sem arquivo em dados/cmed'); return
    edicoes = [(arqs[-1].stem[-8:], ler_planilha(arqs[-1]))]
    for h in sorted((here / 'dados' / 'cmed' / 'historico').glob('*.json.gz'), reverse=True):  # mais nova primeiro
        edicoes.append((h.name[:8], json.loads(gzip.decompress(h.read_bytes()))))
    idx = {k: {} for k in 'prod lab cls sub tipo'.split()}
    ix = lambda k, v: idx[k].setdefault(str(v or '').strip(), len(idx[k]))
    rows, vistos, n_hist = [], set(), 0
    for k, (data_, linhas) in enumerate(edicoes):
        for eans, prod, apres, lab, cls, sub, tipo in linhas:
            eans = [e for e in eans if e not in vistos]
            if not eans or not prod: continue
            vistos.update(eans)
            rows.append([','.join(eans), ix('prod', prod), apres, ix('lab', lab), ix('cls', cls), ix('sub', sub), ix('tipo', tipo)] + ([data_] if k else []))
            n_hist += 1 if k else 0
    inv = lambda d: [k for k, _ in sorted(d.items(), key=lambda kv: kv[1])]
    out = {'fonte': arqs[-1].name + (f' + {len(edicoes) - 1} edições anteriores' if len(edicoes) > 1 else ''), **{k: inv(v) for k, v in idx.items()}, 'rows': rows}
    # substância -> classe mais comum na CMED (para os registrados sem preço)
    sc = defaultdict(Counter)
    for r in rows: sc[';'.join(sorted(N(x) for x in re.split(r'[;,]', out['sub'][r[5]]) if N(x)))][out['cls'][r[4]]] += 1
    sub_cls = {k: v.most_common(1)[0][0] for k, v in sc.items()}
    out['reg'] = registrados({N(p) for p in out['prod']}, sub_cls)
    out['dcb'], out['sais'] = dcb(sub_cls)
    raw = json.dumps(out, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    (here / 'cmed.json.gz').write_bytes(gzip.compress(raw, 9))
    print(f"{arqs[-1].name} + {len(edicoes) - 1} anteriores: {len(rows)} apresentações ({n_hist} só em edições antigas), {len(vistos)} EANs, {len(idx['prod'])} produtos; "
          f"DCB {len(out['dcb'])} substâncias com classe e {len(out['sais'])} sais; "
          f"{len(out['reg'])} registrados fora da lista de preços ({sum(1 for r in out['reg'] if r[2])} com classe pela substância); "
          f"{len(raw)/1e6:.1f} MB -> {(here / 'cmed.json.gz').stat().st_size/1e6:.2f} MB gzip")


if __name__ == '__main__':
    main()
