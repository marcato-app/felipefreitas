import openpyxl,re,collections,sys
S='/tmp/claude-0/-home-user-felipefreitas/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/scratchpad'
U=list(openpyxl.load_workbook('/root/.claude/uploads/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/c66ebe5c-02102026_GABRIEL_COMPLEMENTAR_CHARS_VIT.xlsx',read_only=True).worksheets[0].iter_rows(values_only=True))
N=list(openpyxl.load_workbook(S+'/vit_novo.xlsx',read_only=True).worksheets[0].iter_rows(values_only=True))
nb={}
for r in N[1:]: nb.setdefault(str(r[0]),r)
JUNK=set('HOMEM MULHER KIDS INFANTIL SENIOR GESTANTE ADULTO CABELO CABELOS PELE UNHA UNHAS IMUNIDADE IMUNE VEGANO VEGANA VEGAN COMPLEXO EXTRATO SKIN HAIR NAILS BEAUTY SUP SUPLEMEN ALIMENTA VIT VITA VITAMINA MULT SUPL ALIM OLEO GOMAS'.split())
import json,gzip,csv,glob
_c=json.load(open('consts.json')); _d=json.loads(gzip.decompress(open('dima.json.gz','rb').read()))
FABS_BASE=set(_c['farmaFabs'].values())|set(_d['fabs'])|{m[2] for m in _c['farmaMarcas'] if m[2]}
for _p in glob.glob('dados/vitaminas/*.tsv'):
  for _b in csv.DictReader(open(_p),delimiter='\t'): FABS_BASE.add(_b['Fabricante'].strip())
def sigla_desc(d): m=re.search(r'\(([A-Z0-9&]{3})\)\s*$',str(d or '')); return m.group(1) if m else None
def problema(r):
  m=str(r[3] or ''); s=sigla_desc(r[1]); f=str(r[2] or '')
  core=m[:-4] if s and m.endswith(' '+s) else m
  if core in JUNK: return 'marca e palavra descritiva'
  if m=='OUTRA MARCA' and r[5]!='CODIGO INTERNO': return 'OUTRA MARCA'
  if f=='OUTRO FABRICANTE' and m!='OUTRA MARCA' and not m.endswith(' OTF') and r[5]!='CODIGO INTERNO': return 'marca de outro fabricante'
  return None
out=[]; mud=collections.Counter(); ex=[]
for r in U[1:]:
  p=problema(r); n=nb.get(str(r[0]))
  if p and n:
    novo=list(r); nota=f'CORRIGIDO: {p}'
    so=sigla_desc(r[1]); sn=sigla_desc(n[1]); mn=str(n[3] or '')
    core=mn[:-4] if sn and mn.endswith(' '+sn) else mn
    fo=str(r[2] or ''); so_core=str(r[3] or '')[:-4] if so and str(r[3] or '').endswith(' '+so) else str(r[3] or '')
    inventado = fo not in FABS_BASE and (fo in JUNK or fo==so_core)  # HOMEM/HOMEM, COMPLEXO/COMPLEXO: fabricante inventado pela marca
    fab_ok = fo not in ('','OUTRO FABRICANTE') and not inventado
    if inventado: nota+=f' | fabricante {fo} nao existe nas bases (era a propria marca)'
    if core in JUNK or core=='' or (mn=='OUTRA MARCA' and str(r[3])!='OUTRA MARCA'):
      mud[p+' (sem melhora)']+=1; out.append(list(r)+['']); continue
    if not fab_ok or n[2]==r[2]:  # sem fabricante, ou o mesmo fabricante: tudo novo (marca, fabricante, descritivo, status)
      novo[2],novo[3],novo[8]=n[2],n[3],n[8]
      if n[5]==r[5]: novo[1]=n[1]
    else:  # fabricante do arquivo fica: só aceita marca que é nome do produto (VITAMINA E, ZINCO); marca de outro fabricante não mistura
      PROD=('VITAMINA','VIT','FERRO','CALCIO','ZINCO','MAGNESIO','OMEGA','COLAGENO','BIOTINA','OLEO','CURCUMA','COENZIMA','SELENIO','COMPLEXO','ACIDO','MELATONINA','TRIPTOFANO','CROMO','LUTEINA','VINAGRE','MACA','CHA','FIBRA','CAFEINA','PROPOLIS','GELEIA','LEVEDO','SPIRULINA','CLORETO','DOLOMITA','POLIVITAMINICO','MULTIVITAMINICO','HOMEM','MULHER')
      if core.split()[0] not in PROD:
        mud[p+' (sugestao de outro fabricante, nao aplicada)']+=1; out.append(list(r)+[f'REVISAR: app sugere marca {mn} / fabricante {n[2]}']); continue
      novo[3]=(core+' '+so) if so else core
    if (novo[3],novo[2],novo[1])!=(r[3],r[2],r[1]):
      mud[p]+=1; ex.append((p,r[1],r[3],r[2],'->',novo[1],novo[3],novo[2])); out.append(novo+[nota])
    else: mud[p+' (sem melhora)']+=1; out.append(list(r)+[''])
  else: out.append(list(r)+[''])
print(mud)
for e in ex[::max(1,len(ex)//45)] if len(sys.argv)<2 else []: print(e[0][:12],'|',str(e[1])[:42],'|',e[2],'|',str(e[3])[:16],'->',str(e[5])[:42],'|',e[6],'|',str(e[7])[:18])
if len(sys.argv)>1:
  from openpyxl.styles import Font,PatternFill
  wb=openpyxl.Workbook(); ws=wb.active; ws.title='Planilha1'; ws.append(list(U[0])+['AJUSTE'])
  for x in ws[1]: x.font=Font(bold=True)
  for r in out: ws.append(r)
  for row in ws.iter_rows(min_row=2):
    row[0].number_format='0'
    if row[10].value: 
      for c in row: c.fill=PatternFill('solid',fgColor='FFF2CC')
  ws.auto_filter.ref=ws.dimensions; ws.freeze_panes='B2'
  wb.save(sys.argv[1]); print('salvo',sys.argv[1],len(out))
