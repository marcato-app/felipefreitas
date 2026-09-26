"""Planilha da categoria VITAMINA E MINERAL com o descritivo corrigido pelo app (saída do rodar_vit.js)."""
import csv, json, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
here = Path(__file__).parent
base = list(csv.DictReader(open(here.parent / 'dados' / 'vitaminas' / 'vitamina_mineral.tsv', encoding='utf-8'), delimiter='\t'))
out = json.load(open(here / 'vitaminas_saida.json'))
N = lambda s: re.sub(r'\s+', ' ', re.sub(r'[^A-Z0-9 ]', ' ', str(s or '').upper())).strip()
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Descritivo corrigido'
cab = ['CODIGO BARRAS', 'DESCRICAO ATUAL', 'DESCRITIVO NOVO', 'EST MER 7 ATUAL', 'EST MER 7 APP', 'MARCA ATUAL', 'MARCA APP',
       'FABRICANTE ATUAL', 'FABRICANTE APP', 'CONTEUDO ATUAL', 'CONTEUDO APP', 'CONFERIR']
ws.append(cab)
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864')
n = {'desc': 0, 'conf': 0}
for b, (bc, d, cat, marca, fab, desc, qt, *resto) in zip(base, out):
    reg = bool(resto and resto[0]); nut = resto[1] if len(resto) > 1 else None
    e7 = b['Est Mer 7 Descripcion']; conf = []
    if reg and not N(cat).endswith(N(e7)): conf.append('remédio registrado: categoria de remédio')
    elif not N(cat).endswith(N(e7)) and {cat, e7} == {'VITAMINA OUTRO', 'MULTIVITAMINICO'} and nut not in (None, 0):
        conf.append('regra de nutrientes (até 2 = OUTRO; 3+ ou A-Z = MULTI): conferir a base')
    elif not N(cat).endswith(N(e7)): conf.append('categoria')
    try: q = float(b['Contenido'])
    except ValueError: q = 0
    if q > 1 and abs((qt or 0) - q) > 0.01: conf.append('conteúdo')
    if not desc.startswith('SUPL ALIM') and not reg: conf.append('padrão SUPL ALIM')
    n['desc'] += desc != b['Descripcion']; n['conf'] += bool(conf)
    ws.append([bc, b['Descripcion'], desc, e7, cat, b['Marca'], marca, b['Fabricante'], fab, b['Contenido'], qt, ', '.join(conf)])
for col, w in zip('ABCDEFGHIJKL', [15, 50, 55, 18, 30, 22, 22, 26, 26, 10, 10, 26]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
dst = here.parent / 'dados' / 'vitaminas' / 'vitamina_mineral_corrigido.xlsx'; wb.save(dst)
print(f"{len(base)} SKUs, {n['desc']} com descritivo diferente do atual, {n['conf']} marcados para conferir -> {dst}")
