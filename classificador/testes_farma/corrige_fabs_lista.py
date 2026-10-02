# Corrige, no arquivo GABRIEL já corrigido, as linhas cujo fabricante caiu como "não criar/verificar"
# na comparação da lista de fabricantes com o Hoja 1_46. Decisão linha a linha, a partir do descritivo do PDV.
import csv, json, re, sys, collections, openpyxl
from openpyxl.styles import PatternFill
S = '/tmp/claude-0/-home-user-felipefreitas/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/scratchpad'
BASE = '/home/user/felipefreitas/classificador/'
OF = 'OUTRO FABRICANTE'
CI = 'CI'
# linha do Excel -> (fabricante, marca, observação) | (CI, observação) | (None, observação) = manter
D = {}
def put(rows, v):
    for r in rows: D[r] = v
NV = (None, 'OK: NUTRA VITTON pode ser criado (no Hoja so existe a marca, sem fabricante; NUTRA FOODS e outra empresa)')
put([74, 75, 104, 105, 424, 426, 427, 429, 430, 433, 489, 490, 491, 492, 542, 543, 5233, 5234, 7772], NV)
put([132], ('FORTLIFE', 'BERBERINA FRT', 'PDV: MTC FORTLIFE; FORTLIFE ja existe no Hoja'))
put([5992], ('PERFEITA ALQUIMIA PROD. NAT. LTDA.', 'BERBERINA PER', 'PDV: PERFEITA ALQUIMIA (ja existe no Hoja)'))
put([5988], ('PERFEITA ALQUIMIA PROD. NAT. LTDA.', 'ORA PRO NOBIS PER', 'PDV: PERFEITA ALQUIMIA (ja existe no Hoja)'))
put([249], ('SUPLESSENT', 'SUPLESSENT SUP', 'PDV: SUPLESSENT HOMEM (SUPLESSENT esta na sua lista para criar)'))
put([414, 476, 498, 1324, 3831, 5451, 7972, 7976, 7977, 7978, 1201],
    ('GOMES SUPLEMENTOS ALIMENTARES LTDA', 'GOOD VIT VIT', 'GOOD VIT e marca da GOMES SUPLEMENTOS no Hoja (69 SKUs); nao criar GOODVIT/GOD VIT'))
put([447], ('WOD NUTRITION', 'WOD NUTRITION WOD', 'WODNUTRITION = WOD NUTRITION: criar so um'))
put([501, 11607, 11608], (None, 'OK: WOD NUTRITION pode ser criado (WOW NUTRITION e outra empresa); criar so este, nao o WODNUTRITION'))
put([459], ('NXT', 'CALMIN NXT', 'PDV: NXT CALMIN DIA; CALMIN e produto da NXT'))
put([645], (OF, 'VITAMINA D3 OTF', 'FOUR nao e fabricante; PDV cita NEW FOUR, que nao existe no Hoja'))
put([667], (OF, 'MAGNESIO OTF', 'DEFENSE nao e fabricante; PDV so diz MAGNESIO PA'))
put([716], ('BICAFE', 'BICAFE', 'NAO E FARMA: capsula de cafe BICAFE (fabricante ja existe no Hoja) - mudar categoria'))
put([718], (OF, 'TRIBULUS OTF', 'TRIBULLUS e ingrediente; PDV cita BULL PHARMA, que nao existe no Hoja'))
put([721], ('BIO ERVAS ALIMENTOS', 'BIO ERVAS BIE', 'PDV: EXTRATO DE ERVAS BIO ERVAS (ja existe no Hoja)'))
put([729], ('NUTRILIFE', 'LEVEDO DE CERVEJA NRR', 'PDV: LEVEDO CERVEJA NUTRILIFE'))
put([3858], (OF, 'LEVEDO DE CERVEJA OTF', 'PDV cita UBIMAX, que no Hoja nao tem fabricante'))
put([4368], ('BENNU INDUSTRIA E COMERCIO DE ALIMENTOS', 'MULTINATURAL BEN', 'MULTINATURAL e marca da BENNU no Hoja'))
put([5064], ('MEDIERVAS INDUSTRIA', 'LEVEDO DE CERVEJA RVS', 'PDV: MEDIERVAS'))
put([8664], (OF, 'LEVEDO DE CERVEJA OTF', 'PDV cita LEV LIFE, que nao existe no Hoja'))
put([749], ('P&G', 'ORAL B', 'NAO E FARMA: fita de clareamento dental ORAL-B (P&G) - mudar categoria'))
put([750], ('CHEIRO VERDE', 'CHEIRO VERDE', 'NAO E VITAMINA: cha BAIXA BARRIGA da CHEIRO VERDE (ja existe no Hoja) - conferir categoria'))
put([764], (OF, 'OLEO DE SEMENTE DE ABOBORA OTF', 'SUPLEMEMTO e erro de digitacao, nao fabricante'))
put([3525], (OF, 'OLEO DE SEMENTE DE ABOBORA OTF', 'OL SEMENTE ABOBORA e o produto, nao fabricante'))
put([776], ('KI DELICIA IND E COM ALIM', 'VITAL ERVAS PAM', 'PDV: ORODIS MIX VITAL ERVAS; VITAL ERVAS e marca da KI DELICIA no Hoja - conferir'))
put([4328], (OF, 'PORRETA OTF', 'Energetico PORRETA (no Hoja sem fabricante); conferir se e bebida e nao suplemento'))
put([4330], (CI, 'Produtos distintos no codigo: SUCUPIRA 500MG x ENERGY CALCIO'))
put([4396], (CI, 'Produtos distintos no codigo: ENERGETICO NINJA 10ML x ENERGETICO BALY 250ML'))
put([8543], (CI, 'Produtos distintos no codigo: ENERG TNT JUICE MANGO x COMPOSTO ENERGETICO GUANATUS'))
put([8570], (OF, 'MAIS MIX OTF', 'Fabricante nao identificado'))
put([8947], (OF, 'SAUL OTF', 'Kit energetico SAUL; fabricante nao identificado'))
put([9201], (OF, 'ENERGY CALCIO OTF', 'Fabricante nao identificado (nao e da ENERGY UP)'))
put([8989], (None, 'OK: ENERGY UP ja existe no dicionario (sigla ENR)'))
put([813], (None, 'OK: RENOV pode ser criado (RENO/RENOVIT sao outras empresas)'))
put([1501], (OF, 'SLEEP SHOT OTF', 'SLEEP nao e fabricante'))
put([1857], (None, 'OK: PREVEMAG pode ser criado (PREVEMAX e outra empresa)'))
put([3207], (OF, 'VITAMINA OLEOSA OTF', 'OLEOSA nao e fabricante'))
put([3796], ('SANTA BARBARA', 'SANTA BARBARA SAN', 'STA BARBARA = SANTA BARBARA (ja existe no Hoja, faz chas e temperos) - conferir'))
put([4145], (OF, 'MAGNESIO OTF', 'MAGNESIOS nao e fabricante; PDV: 9 MAGNESIOS ... GREEN'))
put([10101], (OF, 'MAGNESIO OTF', 'MAGNESIOS nao e fabricante; PDV cita MINALAN, que nao existe no Hoja'))
put([4162], ('NUTERAL', 'REABILIT NUT', 'NUTERAL pode ser criado (NUTRAL e outra); a marca do produto e REABILIT'))
put([4225], ('REDE ECONOMIZE', 'SEAME RED', 'SEAME e marca da REDE ECONOMIZE no Hoja; ACAO e de TRIPLA ACAO'))
put([4700], (OF, 'MAGNESIO OTF', 'ACAO e de TRIPLA ACAO, nao fabricante'))
put([9726], ('VIDORA', 'BIOVITA VDO', 'PDV: BIOVITA ... (VIDORA); ACAO e de TRIPLA ACAO'))
put([10516], (OF, 'MINERAZ OTF', 'ACAO e de TRIPLA ACAO; MINERAZ sem fabricante no Hoja'))
put([10969], (OF, 'RESLIV OTF', 'ACAO e de TRIPLA ACAO; RESLIV sem fabricante no Hoja'))
put([4683], (OF, 'ORTO CONDRO OTF', 'ORTO e pedaco do nome ORTO CONDRO'))
put([5603], (OF, 'ORTO CALCIO OTF', 'NAO E VITAMINA: pomada massageadora - conferir categoria'))
put([9279], (OF, 'ORTOCALCIO OTF', 'PDV cita NATURES PRIME, que nao existe no Hoja'))
put([4705], ('FITOTEMP', 'FITOTEMP FIT', 'PRODUTOS NATURAIS nao e fabricante; FITOTEMP ja existe no Hoja'))
put([4735], (CI, 'Produtos distintos no codigo: MOLHO BARBECUE x SUPERERVAS SECA BARRIGA'))
put([4744], (OF, 'AMARGO SLIM OTF', 'AMARGO nao e fabricante'))
put([11398], (OF, 'MULTI EXTRATO OTF', 'AMARGO nao e fabricante'))
put([4831, 6472], (OF, 'TRIBULUS TERRESTRIS OTF', 'TRIBULUS e ingrediente, nao fabricante'))
put([11516], (OF, 'TRIBULUS TERRESTRIS OTF', 'TRIBULLUS e ingrediente, nao fabricante'))
put([10089], (OF, 'TRIBULUS COM MACA OTF', 'PERUVIAN e de MACA PERUANA, nao fabricante'))
put([5027], (OF, 'BERINGELA OTF', 'BERINGELA e o produto; PDV cita AGENUTRY, que nao existe no Hoja'))
put([5053], (OF, 'FIBER SETE OTF', 'FIBER nao e fabricante'))
put([5055], (OF, 'AGUA DE MELISSA OTF', 'AGUA nao e fabricante'))
put([4562], (CI, 'Produtos distintos no codigo: AMPOLA DERMABEL x AMPOLA NUTRILAN ANTICASPA (e nao e vitamina)'))
put([4592], ('SJT', 'VITAMINA K2 SJT', 'PDV: SJT VITAMINA K2 (SJT esta na sua lista para criar)'))
put([5333], ('UNILIFE', 'VITAMINA K2 UNL', 'PDV: UNILIFE (ja existe no Hoja)'))
put([5303], (OF, 'PILULA DO AMOR OTF', 'PILULA nao e fabricante'))
put([5310], (OF, 'SCHWINN', 'NAO E FARMA: bicicleta ergometrica - mudar categoria'))
put([10544], (CI, 'Produtos distintos no codigo: luva UNICARE x luva UNIGLOVES (e nao e vitamina)'))
put([5533], (None, 'OK: NUTRALIVI pode ser criado (NUTRALI e outra)'))
put([5552], ('PROWIN', 'IMEGAR PWN', 'PDV: PROWIN IMEGAR (PROWIN ja existe)'))
put([5569, 5572, 5580], (None, 'OK: BLISFARMA pode ser criado (BIFARMA e outra)'))
put([5601], ('BEM ESTAR', 'BEM ESTAR', 'ANTIBIOTICOS nao e fabricante; PDV: FARINHA DE MACA PERUANA BEM ESTAR - conferir categoria (farinha)'))
put([5976, 5980], (None, 'OK: NUTRALIN pode ser criado (no Hoja so marca sem fabricante)'))
put([9209], ('PRAGERON', 'OLEO DE AVESTRUZ PRA', 'PDV: OLEO DE AVESTRUZ PRAGERON (PRAGERON esta na sua lista para criar)'))
put([9103], (OF, 'OMEGA 3 6 9 OTF', 'OMEGAS OL AVESTRUZ e o produto; PDV cita NUTRAMED, que nao existe no Hoja'))
put([6051], ('D1000', 'GONUTRI D10', 'GONUTRI e marca da D1000 no Hoja'))
put([6451], (OF, 'VITA Z OTF', 'SUMPLEM e SUPLEMENTO abreviado'))
put([6452], (OF, 'FAT SLIM OTF', 'FAT e de FAT SLIM'))
put([6584], (OF, 'SUPRAVIDA LACT OTF', 'LACT nao e fabricante; PDV: SUPRAVIDA (MATHERLLY), sem fabricante no Hoja'))
put([6615, 6616, 6621, 8826, 11315, 11337, 11338, 11339, 11342, 11387],
    ('AROMA BEM ESTAR IND E COM DE PROD NATURAIS', 'SOMALIFE ARO', 'SOMALIFE e marca da AROMA BEM ESTAR no Hoja (sigla ARO)'))
put([6653], (OF, 'CURCUMA OTF', 'JIANG HUANG = curcuma (medicina chinesa), nao fabricante'))
put([11519], (OF, 'CURCUMA OTF', 'ACAFRAO = curcuma, nao fabricante'))
put([6656], (OF, 'CABELOS PELE E UNHAS OTF', 'CABELO nao e fabricante'))
put([6666], (OF, 'ESTICA OTF', 'ESTICA sem fabricante identificado'))
put([6706], (OF, 'POLIVITAMINICO A Z OTF', 'MINERAIS nao e fabricante'))
put([6707], (OF, 'POLIVITAMINICO OTF', 'MUHER (MULHER) nao e fabricante'))
put([7016, 7022, 7025], ('HEALTH LABS', 'HEALTH LABS HEL', 'VITAMINC nao e fabricante; PDV: HEALTH LABS (ja existe no Hoja)'))
put([7045, 7046, 7071, 7110, 7111, 7115, 7116, 7117, 7118, 7129, 7142, 7149], (None, 'OK: ZOMIX pode ser criado (OMIX e outra)'))
put([7143], (None, 'OK: BIOFITHUS ja existe no dicionario (sigla BIH)'))
put([7383], (OF, 'SUPLEMENTAL PLUS OTF', 'SUPLEMENTAL sem fabricante identificado'))
put([7514], ('ENDOGEN', 'COENZIMA Q10 END', 'PDV: COENZYME Q10 PURE DROPS ENDOGEN (ja existe no Hoja)'))
put([8068], (None, 'OK: VITT pode ser criado (VITTO e outra)'))
put([8169], ('LABOGAN', 'SANAR LAB', 'PDV: SANAR; SANAR e marca da LABOGAN no Hoja'))
put([8463], (None, 'OK: HUMAVITA pode ser criado (HUMAITA e outra)'))
put([11057], (None, 'OK: HEALTHDAY pode ser criado (HEALTHY DO BRASIL e outra)'))
put([11436], (None, 'OK: NUTRIGENES pode ser criado (NUTRIGEN e outra)'))
put([11482], (None, 'OK: VITTALINE pode ser criado (VITTALIFE e outra)'))
put([9921, 9922, 10108], (None, 'OK: NUTRIONE pode ser criado (NUTRIZON e outra)'))
put([9480, 9482, 9500], (None, 'OK: VITCORP pode ser criado (VITALCORP e outra)'))
put([9295], (None, 'OK: SELECT NUTRI pode ser criado (SELECT/SELECT GREEN sao outras)'))
put([9296, 9297], ('SELECT NUTRI', 'SELECTNUTRI SEL', 'Marca padronizada para SELECTNUTRI (igual ao outro item do fabricante)'))
put([8690], ('ZYDUS BRASIL', 'CALCIO MAGNESIO ZYD', 'ZYDUS NIKKHO e marca da ZYDUS BRASIL no Hoja (sigla ZYD)'))
put([8805], ('ZYDUS BRASIL', 'MAGNESIO ZYD', 'ZYDUS NIKKHO e marca da ZYDUS BRASIL no Hoja (sigla ZYD)'))
put([8808], ('HERTZ FARMACEUTICA', 'SUPRAVIT KLE', 'PDV: SUPRAVIT; SUPRAVIT e marca da HERTZ no Hoja - conferir se e o mesmo produto'))
put([8833, 11376], ('FARMADI', 'FARMADI FAR', 'FARMA e de FARMADI (FARMADI esta na sua lista para criar)'))
put([9092, 9093, 9094], (OF, 'VITAMINA D3 K2 OTF', 'DIAS e de "60 dias cada pote", nao fabricante'))
put([9900], (OF, 'LITHOLEXAL OTF', 'JOINT nao e fabricante'))
put([9987], ('FPP FRANCHISING LTDA', 'VITALITABS FPP', 'PDV: VITALITABS, marca da FPP FRANCHISING no Hoja'))
put([10016], ('A2F LABORATORIO', 'GRA LABS A2F', 'PDV: GRA LABS, marca da A2F no Hoja'))
put([10075], (OF, 'RAPID SLEEP OTF', 'RAPID e de RAPID SLEEP; PDV cita FUSE LABS, que nao existe no Hoja'))
put([10654], (OF, 'PAZ E AMOR KIDS OTF', 'MODERACAO nao e fabricante'))
put([11504], (OF, 'CALMING SHOT OTF', 'CALMING nao e fabricante'))
put([11563], ('ALQUIMIA', 'ALQUIMIA DA SAUDE ALQ', 'PDV: DETOX SUPERGREENS ALQUIMIA DA SAUDE (marca da ALQUIMIA no Hoja)'))
NAO_FARMA = {716, 749, 5310}

# Hoja: fabricantes e marcas que já existem, para o STATUS
hf, hm = set(), set()
for x in csv.reader(open(S + '/hoja46/Hoja 1 (46).csv', encoding='utf-16'), delimiter='\t'):
    if len(x) > 1: hf.add(x[0].strip()); hm.add(x[1].strip())
c = json.load(open(BASE + 'consts.json'))
hf |= set(c['farmaFabs'].values()); hm |= set(c['farmaSiglas']['marcas'])

src = BASE + 'dados/GABRIEL_COMPLEMENTAR_CHARS_VIT_corrigido.xlsx'
wb = openpyxl.load_workbook(src); ws = wb.active
H = [x.value for x in ws[1]]
iF, iM, iD, iC, iE, iI, iS, iP, iA = (H.index(k) + 1 for k in
    ('FABRICANTE', 'MARCA', 'DESCRITIVO PADRONIZADO', 'CATEGORIA', 'EST MER 7 CODIGO', 'CODIGO INTERNO', 'STATUS', 'PRODUTOS DISTINTOS NO CODIGO', 'AJUSTE'))
ws.cell(1, len(H) + 1, 'AJUSTE FABRICANTES')
iA2 = len(H) + 1
# marcas que ficam no arquivo (linhas não tocadas) também já existem
tocadas = set(D)
for r in range(2, ws.max_row + 1):
    if r not in tocadas and ws.cell(r, iM).value: hm.add(ws.cell(r, iM).value)
amarelo = PatternFill('solid', fgColor='FFF2CC'); verde = PatternFill('solid', fgColor='C6EFCE'); azul = PatternFill('solid', fgColor='DDEBF7')
cnt = collections.Counter()
def sig(m): return m.rsplit(' ', 1)[1] if m and ' ' in m else ''
for r, v in sorted(D.items()):
    fab0, mar0, desc = ws.cell(r, iF).value, ws.cell(r, iM).value, ws.cell(r, iD).value or ''
    if v[0] is None:
        ws.cell(r, iA2, v[1]); ws.cell(r, iA2).fill = verde; cnt['mantido (pode criar)'] += 1; continue
    if v[0] == CI:
        ws.cell(r, iF, OF); ws.cell(r, iM, 'OUTRA MARCA'); ws.cell(r, iC, 'CODIGO INTERNO'); ws.cell(r, iE, '.3.10.1.1.3')
        ws.cell(r, iI, 'Possível Código Interno'); ws.cell(r, iS).value = None; ws.cell(r, iP, 'SIM: ' + v[1].split(': ', 1)[1])
        ws.cell(r, iD, re.sub(r'\s*\([A-Z0-9]{3}\)\s*$', '', desc))
        ws.cell(r, iA2, 'CORRIGIDO: ' + v[1]); cnt['codigo interno'] += 1
    else:
        fab, mar, obs = v
        if mar is None: mar = mar0
        # descritivo: troca o nome da marca antiga e a sigla do fim
        core0 = mar0.rsplit(' ', 1)[0] if mar0 and ' ' in mar0 else (mar0 or '')
        core1 = mar.rsplit(' ', 1)[0] if sig(mar) and ' ' in mar and len(sig(mar)) == 3 else mar
        nd = desc
        if core0 and core0 != core1 and (' ' + core1 + ' ') not in (' ' + nd + ' '):
            nd = nd.replace(' ' + core0 + ' ', ' ' + core1 + ' ', 1) if (' ' + core0 + ' ') in nd else nd
        s1 = sig(mar) if len(sig(mar)) == 3 else ''
        nd = re.sub(r'\s*\([A-Z0-9]{3}\)\s*$', '', nd) + (' (' + s1 + ')' if s1 and r not in NAO_FARMA else '')
        ws.cell(r, iF, fab); ws.cell(r, iM, mar); ws.cell(r, iD, nd)
        st = []
        if mar not in hm: st.append('CRIAR MARCA')
        if fab == OF: st.append('SEM FABRICANTE')
        elif fab not in hf: st.append('CRIAR FABRICANTE')
        ws.cell(r, iS).value = (' E '.join(st).replace('CRIAR MARCA E CRIAR FABRICANTE', 'CRIAR MARCA E FABRICANTE').replace('CRIAR MARCA E SEM', 'CRIAR MARCA; SEM') or None)
        ws.cell(r, iA2, 'CORRIGIDO: ' + obs + ' (antes: %s / %s)' % (fab0, mar0))
        cnt['nao e farma' if r in NAO_FARMA else ('OTF' if fab == OF else 'fabricante real')] += 1
    for col in range(1, iA2 + 1):
        ws.cell(r, col).fill = azul if r in NAO_FARMA or v[0] == CI else amarelo
ws.column_dimensions[openpyxl.utils.get_column_letter(iA2)].width = 90
out = BASE + 'dados/GABRIEL_COMPLEMENTAR_CHARS_VIT_corrigido_v2.xlsx'
wb.save(out); print(cnt, len(D), out)
