import openpyxl,sys,json,collections
S='/tmp/claude-0/-home-user-felipefreitas/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/scratchpad'
sys.path.insert(0,S); from veredito import V
from openpyxl.styles import Font,PatternFill,Alignment
rows=list(openpyxl.load_workbook(S+'/vm500s_out.xlsx',read_only=True).worksheets[0].iter_rows(values_only=True)); h=list(rows[0])
src=list(openpyxl.load_workbook('/root/.claude/uploads/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/f8cbb47b-products_-_2026-10-01T235645.827.xlsx',read_only=True).worksheets[0].iter_rows(values_only=True))[1:]
o={str(r[0]):list(r) for r in rows[1:]}
H=h+['E FARMA?','CATEGORIA OUTRO APP']
iC,iM,iF,iE7,iST,iD=h.index('CATEGORIA'),h.index('MARCA'),h.index('FABRICANTE'),h.index('EST MER 7 CODIGO'),h.index('STATUS'),h.index('PRODUTOS DISTINTOS NO CODIGO')
out=[]
for i,s in enumerate(src):
  r=o[str(s[1])][:]; f,cat,obs=V[i]
  if 'CONFLITO' in obs and r[iC]!='CODIGO INTERNO':
    r[iD]='SIM: '+obs.replace('CONFLITO: ',''); r[iC]='CODIGO INTERNO'; r[iM]='OUTRA MARCA'; r[iF]='OUTRO FABRICANTE'; r[iST]=None
  if r[iC]=='CODIGO INTERNO': r[iE7]=None; r[iST]=None
  if r[iD] and str(r[iD]).startswith('VERIFICAR'): r[iD]=None  # mesma cesta: mesmo produto escrito de outro jeito (revisado)
  farma='CODIGO INTERNO' if r[iC]=='CODIGO INTERNO' else ('SIM' if f=='S' else 'NAO')
  out.append(r+[farma,s[0]])
ordem={'SIM':0,'NAO':1,'CODIGO INTERNO':2}
out.sort(key=lambda r:(ordem[r[-2]],str(r[iC]),str(r[1])))
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Planilha1'; ws.append(H)
for x in ws[1]: x.font=Font(bold=True,color='FFFFFF'); x.fill=PatternFill('solid',fgColor='1F3864'); x.alignment=Alignment(wrap_text=True,vertical='top')
for r in out: ws.append(r)
for col,w in zip('ABCDEFGHIJKL',[16,60,28,24,11,40,18,22,26,50,14,24]): ws.column_dimensions[col].width=w
cor={'SIM':'C6EFCE','NAO':'FFC7CE','CODIGO INTERNO':'FFEB9C'}
for row in ws.iter_rows(min_row=2):
  row[0].number_format='0'; row[10].fill=PatternFill('solid',fgColor=cor[row[10].value])
  if row[8].value: row[8].fill=PatternFill('solid',fgColor='FCE4D6')
ws.auto_filter.ref=ws.dimensions; ws.freeze_panes='B2'
p='dados/vitamina_mineral_500_classificado.xlsx'; wb.save(p)
c=collections.Counter(r[-2] for r in out); print(c, collections.Counter(r[iST] for r in out))
print('CODIGO INTERNO:',sum(1 for r in out if r[iC]=='CODIGO INTERNO'))
for r in out[:3]+out[300:303]+out[-3:]: print(r[:10])
