S='/tmp/claude-0/-home-user-felipefreitas/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/scratchpad'
import sys, json, collections, re, difflib, openpyxl
from openpyxl.styles import PatternFill, Font
exec(open(S+'/fab_similar.py').read().split("if __name__")[0])
c=json.load(open('/home/user/felipefreitas/classificador/consts.json'))
dimaMarca={m[0] for m in c['farmaDima']}
SEM_DONO={'OUTRO FABRICANTE','SIN PROVEEDOR ASOCIADO','OUTRA MARCA','POR DEFECTO','Por Defecto',''}
marcaFab=collections.defaultdict(collections.Counter)
for x in r[1:]:
    if x and x[1].strip(): marcaFab[x[1].strip()][x[0].strip()]+=int(x[8]) if x[8].strip().isdigit() else 1
ns={}
for f in fabs: ns.setdefault(cores[f].replace(' ',''),[]).append(f)
# descreve o produto, ingrediente ou erro de digitação dessas palavras
DESC=set(('HOMEM MULHER MUHER VEGANO VEGAN COMPLEXO EXTRATO VITAMINA VITAMINC CABELO CABELOS HAIR AGUA LUVAS ANTIBIOTICOS PILULA '
 'BRANQUEAMENTO SUPLEMEMTO SUPLEMENTAL SUMPLEM MINERAIS MAGNESIOS FIBER SLEEP DETOX CALMING JOINT ANTICASPA '
 'BERBERINA TRIBULUS TRIBULLUS ACAFRAO BERINGELA MENAQUINONA K2 OLEOSA AMARGO NUTRACEUTICA DEFENSE FARMA').split())
FRASES={'LEVEDO DE CERVEJA','PRODUTOS NATURAIS','OL SEMENTE ABOBORA','JIANG HUANG','COENZYME Q10 DROPS','OMEGAS OL AVESTRUZ'}
# palavra solta / cortada: provavelmente não é fabricante, mas conferir no item
SOLTA={'DIAS','FOUR','BIKE','ESTICA','ORTO','GRA','RAPID','MODERACAO','BAIXA','FAT','LACT','PERUVIAN'}
lista=[l.strip() for l in open(S+'/lista274.txt') if l.strip()]
norm_l={}
for n in lista: norm_l.setdefault(core(n).replace(' ',''),[]).append(n)
cor={'JA EXISTE NO HOJA':'F8CBAD','NAO E FABRICANTE':'FFC7CE','PALAVRA SOLTA - CONFERIR':'FFD966','E MARCA NO HOJA':'FCE4D6',
     'VERIFICAR':'FFF2CC','DUPLICADO NA LISTA':'DDEBF7','PODE CRIAR':'C6EFCE'}
wb=openpyxl.Workbook(); w=wb.active; w.title='DETALHE'
w.append(['FABRICANTE DA LISTA','RESULTADO','O QUE FAZER','PARECIDO ENCONTRADO NO HOJA','TIPO DA SEMELHANCA','QTD SKU NO HOJA','MARCAS DELE NO HOJA','CATEGORIAS NO HOJA'])
cnt=collections.Counter(); resumo=[]
for n in lista:
    up=n.upper(); linhas=[]; words=set(core(n).split())
    if up in DESC or up in FRASES or (words and words<=DESC):
        tag='NAO E FABRICANTE'; acao='Palavra descritiva/ingrediente: usar o fabricante real do produto (ou OUTRO FABRICANTE)'
    elif up in SOLTA:
        tag='PALAVRA SOLTA - CONFERIR'; acao='Parece palavra cortada da descricao, nao fabricante: conferir no item'
    elif up in info:
        tag='JA EXISTE NO HOJA'; acao='Nao criar: usar o existente'; linhas=[(up,'IGUAL')]
    else:
        hit=[f for f in ns.get(core(n).replace(' ',''),[]) if f!=up]
        if hit: tag='JA EXISTE NO HOJA'; acao='Nao criar: mesmo nome escrito diferente'; linhas=[(f,'MESMO NOME (espaco/LTDA/GRUPO/pontuacao)') for f in hit]
        else:
            b=busca(n); mf=marcaFab.get(up)
            donos=[f for f,_ in mf.most_common(3) if f not in SEM_DONO] if mf else []
            if donos:
                tag='E MARCA NO HOJA'; acao='E marca no Hoja: usar o fabricante dono dela, nao criar fabricante'
                linhas=[(f,'MARCA "%s" PERTENCE A ESTE FABRICANTE'%up) for f in donos]+[(f,t) for t,f in b[:3]]
            elif b:
                tag='VERIFICAR'; acao='Conferir se e o mesmo; se nao for, pode criar'; linhas=[(f,t) for t,f in b[:6]]
                if mf: acao+=' | no Hoja ja existe como MARCA sem fabricante'
            else:
                tag='PODE CRIAR'; acao='Nada parecido no Hoja'
                if mf: acao+=' (no Hoja ja existe como MARCA sem fabricante)'
    dup=[o for o in norm_l.get(core(n).replace(' ',''),[]) if o!=n]
    if not dup:
        dup=[o for o in lista if o!=n and len(o)>=5 and len(n)>=5 and difflib.SequenceMatcher(None,core(o).replace(' ',''),core(n).replace(' ','')).ratio()>=0.88]
    if dup:
        if tag=='PODE CRIAR': tag='DUPLICADO NA LISTA'; acao='Criar um so para: '+', '.join([n]+dup)
        else: acao+=' | na propria lista tambem tem: '+', '.join(dup)
    if up in dimaMarca and tag not in ('NAO E FABRICANTE',): acao+=' | no DIMA isto e MARCA'
    cnt[tag]+=1; resumo.append((n,tag,acao,linhas[0][0] if linhas else ''))
    for f,t in (linhas or [('','')]):
        i=info.get(f) if f else None
        w.append([n,tag,acao,f,t,i[0] if i else '',', '.join(m for m,_ in i[1].most_common(5)) if i else '',', '.join(m for m,_ in i[2].most_common(3)) if i else ''])
        for cc in w[w.max_row][:2]: cc.fill=PatternFill('solid',fgColor=cor[tag])
w2=wb.create_sheet('RESUMO',0); w2.append(['FABRICANTE DA LISTA','RESULTADO','O QUE FAZER','PRINCIPAL PARECIDO NO HOJA'])
ordem=list(cor)
for x in sorted(resumo,key=lambda x:(ordem.index(x[1]),x[0])):
    w2.append(list(x))
    for cc in w2[w2.max_row][:2]: cc.fill=PatternFill('solid',fgColor=cor[x[1]])
for ws,wd in ((w,[30,26,70,42,36,10,50,40]),(w2,[30,26,80,42])):
    for cc in ws[1]: cc.font=Font(bold=True)
    for i,v in enumerate(wd): ws.column_dimensions['ABCDEFGH'[i]].width=v
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
wb.save('/home/user/felipefreitas/classificador/dados/fabricantes_novos_vs_hoja.xlsx')
print(cnt)
for x in resumo:
    if x[1] in ('E MARCA NO HOJA','JA EXISTE NO HOJA','DUPLICADO NA LISTA','VERIFICAR'): print(x[1][:10],'|',x[0],'->',x[3])
