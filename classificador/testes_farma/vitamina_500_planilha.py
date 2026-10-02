import json,sys,openpyxl,re,collections
S='/tmp/claude-0/-home-user-felipefreitas/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/scratchpad'
sys.path.insert(0,S); from veredito import V
from openpyxl.styles import Font,PatternFill,Alignment
out=list(openpyxl.load_workbook(S+'/vm500_out.xlsx',read_only=True).worksheets[0].iter_rows(values_only=True)); h=out[0]
o={str(r[0]):r for r in out[1:]}
src=list(openpyxl.load_workbook('/root/.claude/uploads/7443e5f8-7000-5ee7-ae7b-4c71ad3bdb8e/f8cbb47b-products_-_2026-10-01T235645.827.xlsx',read_only=True).worksheets[0].iter_rows(values_only=True))[1:]
c=json.load(open('consts.json')); of=c['nomesOficiais']; cats=c['libSeed']['categorias']+c['cfg']['categorias']
nomes={of.get(x['nome'],x['nome']) for x in cats}; cesta={of.get(x['nome'],x['nome']):x.get('cesta','') for x in cats}
em7=c.get('estMer7',{}); inv={v:k for k,v in of.items()}
H=['CODIGO BARRAS','DESCRICOES PDV','CATEGORIA OUTRO APP','CATEGORIA CLASSIFICADOR','E FARMA?','CATEGORIA SUGERIDA (REVISAO)','CESTA','OBSERVACAO','PRODUTOS DISTINTOS NO CODIGO','DESCRITIVO PADRONIZADO (CLASSIFICADOR)','MARCA','FABRICANTE','EST MER 7 CODIGO (SUGERIDA)']
linhas=[]
for i,s in enumerate(src):
  r=o.get(str(s[1])); f,cat,obs=V[i]
  ds=re.findall(r'"((?:[^"\\]|\\.)*)"', s[2] or '')
  dist=(r[h.index('PRODUTOS DISTINTOS NO CODIGO')] if r else '') or ''
  if 'CONFLITO' in obs and not dist: dist='SIM (revisao manual)'
  linhas.append([s[1],' | '.join(ds[:4])[:300],s[0],r[5] if r else '','SIM' if f=='S' else 'NAO',cat,cesta.get(cat,'?') or '(sem cesta)',obs,dist,r[1] if r else '',r[3] if r else '',r[2] if r else '',em7.get(inv.get(cat,cat),'')])
print('categorias inexistentes:',collections.Counter(l[5] for l in linhas if l[5] not in nomes))
wb=openpyxl.Workbook(); wb.remove(wb.active)
def aba(nome,rows):
  ws=wb.create_sheet(nome); ws.append(H)
  for x in ws[1]: x.font=Font(bold=True,color='FFFFFF'); x.fill=PatternFill('solid',fgColor='1F3864'); x.alignment=Alignment(wrap_text=True,vertical='top')
  for l in rows: ws.append(l)
  for col,w in zip('ABCDEFGHIJKLM',[16,60,24,30,9,40,16,45,45,50,22,24,16]): ws.column_dimensions[col].width=w
  for row in ws.iter_rows(min_row=2):
    row[0].number_format='0'; row[4].fill=PatternFill('solid',fgColor='C6EFCE' if row[4].value=='SIM' else 'FFC7CE')
    if row[8].value: row[8].fill=PatternFill('solid',fgColor='FFEB9C')
  ws.auto_filter.ref=ws.dimensions; ws.freeze_panes='C2'
fa=[l for l in linhas if l[4]=='SIM']; nf=[l for l in linhas if l[4]=='NAO']
ws=wb.create_sheet('RESUMO')
res=[['VITAMINA E MINERAL - revisao de 500 itens que outro app colocou em outras categorias',''],['Itens',len(linhas)],['E farmacia (SIM)',len(fa)],['Nao e farmacia (NAO)',len(nf)],['Codigo com produtos distintos',sum(1 for l in linhas if l[8])],['Outro app acertou a categoria (igual a sugerida)',sum(1 for l in linhas if l[2]==l[5])],['Classificador acertou se e farma (cesta)',sum(1 for l in linhas if (l[4]=='SIM')==(cesta.get(l[3],'')=='FARMACIA'))],[],['FARMA por categoria sugerida','']]
res+=[[k,v] for k,v in collections.Counter(l[5] for l in fa).most_common()]
res+=[[],['NAO FARMA por categoria sugerida','']]+[[k,v] for k,v in collections.Counter(l[5] for l in nf).most_common()]
for x in res: ws.append(x)
ws.column_dimensions['A'].width=60; ws['A1'].font=Font(bold=True)
aba('FARMA',fa); aba('NAO FARMA',nf); aba('PRODUTOS DISTINTOS',[l for l in linhas if l[8]]); aba('TODOS',linhas)
wb.save('dados/vitamina_mineral_revisao_500.xlsx'); print(res[1:7])
